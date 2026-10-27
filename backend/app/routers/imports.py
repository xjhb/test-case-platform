import csv
import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user, require_project_access
from ..logger import log_op
from ..models import TestCase, User

router = APIRouter(prefix="/api/imports", tags=["用例导入"], dependencies=[Depends(get_current_user)])

COLUMNS = ["模块", "用例标题", "前置条件", "测试步骤", "预期结果", "优先级"]
VALID_PRIORITY = {"P0", "P1", "P2", "P3"}


@router.get("/template")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "用例模板"
    ws.append(COLUMNS)
    ws.append(["登录模块", "用户使用正确账号密码登录成功", "已打开登录页", "1.输入账号\n2.输入密码\n3.点击登录", "跳转到首页", "P1"])
    ws.append(["下单模块", "库存充足时下单成功", "已登录且商品在购物车", "1.进入购物车\n2.勾选商品\n3.点击结算\n4.提交订单", "生成订单并跳转订单页", "P0"])
    widths = [16, 36, 20, 36, 28, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=testcase_template.xlsx"},
    )


@router.post("")
async def import_testcases(
    project_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="请上传文件")
    require_project_access(project_id, db, current)
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    content = await file.read()

    if ext == "xlsx":
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        data = [list(r) for r in ws.iter_rows(values_only=True)]
    elif ext == "csv":
        text = content.decode("utf-8-sig")
        data = list(csv.reader(io.StringIO(text)))
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .csv 文件")

    if not data or len(data) < 2:
        raise HTTPException(status_code=400, detail="文件没有数据行(需包含表头+至少一行数据)")

    created = 0
    failed = 0
    errors = []
    for idx, row in enumerate(data[1:], start=2):
        if not row or all((c is None or str(c).strip() == "") for c in row):
            continue
        cells = [(str(c).strip() if c is not None else "") for c in row]
        while len(cells) < 6:
            cells.append("")
        module, title, precondition, steps, expected, priority = cells[:6]
        if not title:
            failed += 1
            errors.append(f"第{idx}行:用例标题为空,已跳过")
            continue
        if priority not in VALID_PRIORITY:
            priority = "P1"
        db.add(
            TestCase(
                project_id=project_id,
                module=module,
                title=title,
                precondition=precondition,
                steps=steps,
                expected=expected,
                priority=priority,
                status="未执行",
                created_by=current.id,
            )
        )
        created += 1
    db.commit()
    log_op(db, current, "导入", target=f"项目#{project_id}",
           detail=f"文件:{file.filename} 成功{created} 失败{failed}")
    return {"created": created, "failed": failed, "errors": errors}
