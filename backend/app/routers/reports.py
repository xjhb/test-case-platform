import io
import json
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import accessible_project_ids, get_current_user, require_project_access
from ..logger import log_op
from ..models import Defect, Project, TestCase, TestPlan, TestReport, TestRun, User
from ..schemas import TestReportOut

router = APIRouter(prefix="/api/reports", tags=["测试报告"], dependencies=[Depends(get_current_user)])

STAGE_ORDER = ["SIT", "UAT", "PRE"]
STAGE_COLOR = {"通过": "C6EFCE", "失败": "FFC7CE", "阻塞": "FFEB9C", "未执行": "F2F2F2"}
THIN = Side(style="thin", color="D9D9D9")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _latest_run_map(db: Session, plan_id: int, stage: str) -> dict[int, TestRun]:
    """返回 {testcase_id: 该阶段最新的TestRun}"""
    runs = list(
        db.execute(
            select(TestRun)
            .where(TestRun.plan_id == plan_id, TestRun.stage == stage)
            .order_by(TestRun.id.desc())
        ).scalars()
    )
    latest: dict[int, TestRun] = {}
    for r in runs:
        if r.testcase_id not in latest:
            latest[r.testcase_id] = r
    return latest


def _user_name(db: Session, uid: int | None) -> str:
    if not uid:
        return ""
    u = db.get(User, uid)
    return u.display_name or u.username if u else ""


@router.get("/plan/{plan_id}/export")
def export_plan_report(plan_id: int, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    plan = db.get(TestPlan, plan_id)
    if not plan:
        raise HTTPException(status_code=404, detail="计划不存在")
    require_project_access(plan.project_id, db, current)
    project = db.get(Project, plan.project_id)
    cases = list(
        db.execute(select(TestCase).where(TestCase.project_id == plan.project_id).order_by(TestCase.id.asc())).scalars()
    )

    stage_maps = {st: _latest_run_map(db, plan_id, st) for st in STAGE_ORDER}

    wb = Workbook()

    # 公共样式
    title_font = Font(size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="1E3C72")
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")

    # ============ Sheet1: 测试概览 ============
    ws1 = wb.active
    ws1.title = "测试概览"

    ws1.merge_cells("A1:G1")
    ws1["A1"] = "测 试 报 告"
    ws1["A1"].font = title_font
    ws1["A1"].fill = title_fill
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 30

    info = [
        ("项目名称", project.name if project else "", "计划版本", plan.version),
        ("当前阶段", plan.current_stage, "计划状态", plan.status),
        ("用例总数", str(len(cases)), "生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for i, (k1, v1, k2, v2) in enumerate(info, start=2):
        ws1.cell(row=i, column=1, value=k1).font = Font(bold=True)
        ws1.cell(row=i, column=2, value=v1)
        ws1.cell(row=i, column=4, value=k2).font = Font(bold=True)
        ws1.cell(row=i, column=5, value=v2)
        ws1.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        ws1.merge_cells(start_row=i, start_column=5, end_row=i, end_column=7)

    # 阶段统计表
    start = 6
    ws1.cell(row=start, column=1, value="阶段执行统计").font = Font(bold=True, size=12)
    headers = ["阶段", "用例总数", "通过", "失败", "阻塞", "未执行", "通过率"]
    for c, h in enumerate(headers, start=1):
        cell = ws1.cell(row=start + 1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = CELL_BORDER
    total = len(cases)
    for i, st in enumerate(STAGE_ORDER):
        m = stage_maps[st]
        passed = failed = blocked = pending = 0
        for c in cases:
            run = m.get(c.id)
            s = run.status if run else "未执行"
            if s == "通过":
                passed += 1
            elif s == "失败":
                failed += 1
            elif s == "阻塞":
                blocked += 1
            else:
                pending += 1
        rate = f"{passed / total * 100:.1f}%" if total else "0%"
        row = start + 2 + i
        vals = [st, total, passed, failed, blocked, pending, rate]
        for c, v in enumerate(vals, start=1):
            cell = ws1.cell(row=row, column=c, value=v)
            cell.alignment = Alignment(horizontal="center")
            cell.border = CELL_BORDER

    widths1 = [14, 14, 12, 12, 12, 12, 12]
    for i, w in enumerate(widths1, 1):
        ws1.column_dimensions[chr(64 + i)].width = w

    # ============ Sheet2: 用例明细 ============
    ws2 = wb.create_sheet("用例明细")
    cols = [
        "ID", "模块", "用例标题", "优先级", "前置条件", "测试步骤", "预期结果",
        "SIT结果", "SIT执行人", "SIT执行时间", "SIT备注",
        "UAT结果", "UAT执行人", "UAT执行时间", "UAT备注",
        "PRE结果", "PRE执行人", "PRE执行时间", "PRE备注",
        "当前阶段状态",
    ]
    for c, h in enumerate(cols, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER

    # 缓存执行人姓名
    executor_cache: dict[int, str] = {}

    for r, tc in enumerate(cases, start=2):
        ws2.cell(row=r, column=1, value=tc.id).border = CELL_BORDER
        ws2.cell(row=r, column=2, value=tc.module).border = CELL_BORDER
        ws2.cell(row=r, column=3, value=tc.title).border = CELL_BORDER
        ws2.cell(row=r, column=4, value=tc.priority).border = CELL_BORDER
        ws2.cell(row=r, column=5, value=tc.precondition).border = CELL_BORDER
        ws2.cell(row=r, column=6, value=tc.steps).border = CELL_BORDER
        ws2.cell(row=r, column=7, value=tc.expected).border = CELL_BORDER
        col = 8
        for st in STAGE_ORDER:
            run = stage_maps[st].get(tc.id)
            s = run.status if run else "未执行"
            # 结果(带颜色)
            cell = ws2.cell(row=r, column=col, value=s)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor=STAGE_COLOR.get(s, "F2F2F2"))
            cell.border = CELL_BORDER
            # 执行人
            executor = ""
            if run and run.executed_by:
                if run.executed_by not in executor_cache:
                    executor_cache[run.executed_by] = _user_name(db, run.executed_by)
                executor = executor_cache[run.executed_by]
            ws2.cell(row=r, column=col + 1, value=executor).border = CELL_BORDER
            # 执行时间
            ts = run.executed_at.strftime("%Y-%m-%d %H:%M") if run and run.executed_at else ""
            ws2.cell(row=r, column=col + 2, value=ts).border = CELL_BORDER
            # 备注
            ws2.cell(row=r, column=col + 3, value=run.remark if run else "").border = CELL_BORDER
            col += 4
        # 当前阶段状态
        cur_run = stage_maps[plan.current_stage].get(tc.id)
        cur_s = cur_run.status if cur_run else "未执行"
        cur_cell = ws2.cell(row=r, column=20, value=cur_s)
        cur_cell.alignment = Alignment(horizontal="center")
        cur_cell.fill = PatternFill("solid", fgColor=STAGE_COLOR.get(cur_s, "F2F2F2"))
        cur_cell.border = CELL_BORDER

    widths2 = [6, 12, 26, 8, 20, 26, 22,
               10, 12, 18, 20,
               10, 12, 18, 20,
               10, 12, 18, 20,
               14]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = w
    ws2.freeze_panes = "D2"

    # ============ Sheet3: 失败用例汇总 ============
    ws3 = wb.create_sheet("失败阻塞汇总")
    cols3 = ["阶段", "ID", "模块", "用例标题", "结果", "执行人", "执行时间", "备注"]
    for c, h in enumerate(cols3, start=1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = PatternFill("solid", fgColor="C00000")
        cell.alignment = Alignment(horizontal="center")
        cell.border = CELL_BORDER
    row3 = 2
    for st in STAGE_ORDER:
        for tc in cases:
            run = stage_maps[st].get(tc.id)
            if run and run.status in ("失败", "阻塞"):
                vals = [st, tc.id, tc.module, tc.title, run.status,
                        executor_cache.get(run.executed_by, _user_name(db, run.executed_by)),
                        run.executed_at.strftime("%Y-%m-%d %H:%M") if run.executed_at else "",
                        run.remark]
                for c, v in enumerate(vals, start=1):
                    cell = ws3.cell(row=row3, column=c, value=v)
                    cell.border = CELL_BORDER
                    if c == 5:
                        cell.fill = PatternFill("solid", fgColor=STAGE_COLOR.get(v, "F2F2F2"))
                        cell.alignment = Alignment(horizontal="center")
                row3 += 1
    if row3 == 2:
        ws3.cell(row=2, column=1, value="无失败/阻塞记录").alignment = Alignment(horizontal="center")
        ws3.merge_cells("A2:H2")
    widths3 = [10, 8, 12, 26, 10, 12, 18, 24]
    for i, w in enumerate(widths3, 1):
        ws3.column_dimensions[chr(64 + i)].width = w
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_op(db, current, "导出", target=f"测试报告 计划#{plan.id} {plan.version}")
    fname = f"test_report_{plan.version}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ========== 自动生成测试报告(持久化) ==========

def _generate_report_data(db: Session, plan_id: int, stage: str, current: User) -> dict:
    plan = db.get(TestPlan, plan_id)
    if not plan:
        raise HTTPException(404, "测试计划不存在")
    require_project_access(plan.project_id, db, current)

    runs = list(db.execute(
        select(TestRun).where(TestRun.plan_id == plan_id, TestRun.stage == stage)
    ).scalars())

    testcase_ids = {r.testcase_id for r in runs}
    cases = list(db.execute(select(TestCase).where(TestCase.id.in_(testcase_ids))).scalars()) if testcase_ids else []
    case_map = {c.id: c for c in cases}

    total = len(testcase_ids)
    passed = sum(1 for r in runs if r.status == "通过")
    failed = sum(1 for r in runs if r.status == "失败")
    blocked = sum(1 for r in runs if r.status == "阻塞")
    unexecuted = total - passed - failed - blocked
    pass_rate = round(passed / total * 100, 2) if total > 0 else 0.0

    defects = list(db.execute(
        select(Defect).where(Defect.plan_id == plan_id, Defect.status != "关闭")
    ).scalars())
    defect_count = len(defects)
    critical_defects = sum(1 for d in defects if d.severity in ("致命", "严重"))

    summary_data = {
        "plan": {"id": plan.id, "version": plan.version, "project_id": plan.project_id},
        "stage": stage,
        "runs": [
            {
                "testcase_id": r.testcase_id,
                "testcase_title": case_map.get(r.testcase_id, TestCase(title="")).title,
                "status": r.status,
                "remark": r.remark,
                "executed_by": r.executed_by,
                "executed_at": r.executed_at.isoformat() if r.executed_at else "",
            }
            for r in runs
        ],
        "defects": [
            {
                "id": d.id,
                "title": d.title,
                "severity": d.severity,
                "priority": d.priority,
                "status": d.status,
                "assignee_id": d.assignee_id,
            }
            for d in defects
        ],
    }

    return {
        "plan_id": plan_id,
        "stage": stage,
        "title": f"{plan.version} - {stage} 测试报告",
        "total_cases": total,
        "passed_cases": passed,
        "failed_cases": failed,
        "blocked_cases": blocked,
        "unexecuted_cases": unexecuted,
        "pass_rate": pass_rate,
        "defect_count": defect_count,
        "critical_defects": critical_defects,
        "summary": json.dumps(summary_data, ensure_ascii=False),
    }


@router.post("/generate", response_model=TestReportOut)
def generate_report(
    plan_id: int,
    stage: str = "SIT",
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    if stage not in ("SIT", "UAT", "PRE"):
        raise HTTPException(400, "阶段只能是 SIT/UAT/PRE")

    data = _generate_report_data(db, plan_id, stage, current)

    existing = db.execute(
        select(TestReport).where(TestReport.plan_id == plan_id, TestReport.stage == stage)
    ).scalar_one_or_none()

    if existing:
        for k, v in data.items():
            setattr(existing, k, v)
        existing.generated_by = current.id
        report = existing
    else:
        report = TestReport(**data, generated_by=current.id)
        db.add(report)

    db.commit()
    db.refresh(report)

    generator = db.get(User, report.generated_by)
    log_op(db, current, "生成", target=f"测试报告#{report.id}", detail=f"{report.title}")

    return TestReportOut(
        id=report.id, plan_id=report.plan_id, stage=report.stage, title=report.title,
        total_cases=report.total_cases, passed_cases=report.passed_cases,
        failed_cases=report.failed_cases, blocked_cases=report.blocked_cases,
        unexecuted_cases=report.unexecuted_cases, pass_rate=report.pass_rate,
        defect_count=report.defect_count, critical_defects=report.critical_defects,
        generated_by=report.generated_by,
        generated_by_name=generator.display_name if generator else "",
        generated_at=report.generated_at, summary=report.summary,
    )


@router.get("/list", response_model=list[TestReportOut])
def list_reports(
    plan_id: int | None = None,
    project_id: int | None = None,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    q = select(TestReport)
    if plan_id is not None:
        q = q.where(TestReport.plan_id == plan_id)
    if project_id is not None:
        plan_ids = [p.id for p in db.execute(
            select(TestPlan.id).where(TestPlan.project_id == project_id)
        ).scalars()]
        if plan_ids:
            q = q.where(TestReport.plan_id.in_(plan_ids))
        else:
            return []

    reports = list(db.execute(q.order_by(TestReport.id.desc())).scalars())
    result = []
    for r in reports:
        generator = db.get(User, r.generated_by)
        result.append(TestReportOut(
            id=r.id, plan_id=r.plan_id, stage=r.stage, title=r.title,
            total_cases=r.total_cases, passed_cases=r.passed_cases,
            failed_cases=r.failed_cases, blocked_cases=r.blocked_cases,
            unexecuted_cases=r.unexecuted_cases, pass_rate=r.pass_rate,
            defect_count=r.defect_count, critical_defects=r.critical_defects,
            generated_by=r.generated_by,
            generated_by_name=generator.display_name if generator else "",
            generated_at=r.generated_at, summary=r.summary,
        ))
    return result


@router.get("/{report_id}", response_model=TestReportOut)
def get_report(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    r = db.get(TestReport, report_id)
    if not r:
        raise HTTPException(404, "报告不存在")
    plan = db.get(TestPlan, r.plan_id)
    if plan:
        require_project_access(plan.project_id, db, current)
    generator = db.get(User, r.generated_by)
    return TestReportOut(
        id=r.id, plan_id=r.plan_id, stage=r.stage, title=r.title,
        total_cases=r.total_cases, passed_cases=r.passed_cases,
        failed_cases=r.failed_cases, blocked_cases=r.blocked_cases,
        unexecuted_cases=r.unexecuted_cases, pass_rate=r.pass_rate,
        defect_count=r.defect_count, critical_defects=r.critical_defects,
        generated_by=r.generated_by,
        generated_by_name=generator.display_name if generator else "",
        generated_at=r.generated_at, summary=r.summary,
    )


@router.get("/{report_id}/export")
def export_generated_report(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    r = db.get(TestReport, report_id)
    if not r:
        raise HTTPException(404, "报告不存在")
    plan = db.get(TestPlan, r.plan_id)
    if plan:
        require_project_access(plan.project_id, db, current)

    data = json.loads(r.summary) if r.summary else {}
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "概要"
    ws1.merge_cells("A1:D1")
    title_cell = ws1.cell(row=1, column=1, value=r.title)
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center")

    summary_rows = [
        ("测试阶段", r.stage),
        ("用例总数", r.total_cases),
        ("通过", r.passed_cases),
        ("失败", r.failed_cases),
        ("阻塞", r.blocked_cases),
        ("未执行", r.unexecuted_cases),
        ("通过率", f"{r.pass_rate}%"),
        ("缺陷总数", r.defect_count),
        ("严重缺陷", r.critical_defects),
    ]
    for i, (k, v) in enumerate(summary_rows, start=3):
        ws1.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws1.cell(row=i, column=2, value=v)
    ws1.column_dimensions["A"].width = 15
    ws1.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("执行明细")
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    cols2 = ["用例ID", "用例标题", "执行状态", "备注", "执行人ID", "执行时间"]
    for c, h in enumerate(cols2, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for i, run in enumerate(data.get("runs", []), start=2):
        ws2.cell(row=i, column=1, value=run.get("testcase_id", ""))
        ws2.cell(row=i, column=2, value=run.get("testcase_title", ""))
        ws2.cell(row=i, column=3, value=run.get("status", ""))
        ws2.cell(row=i, column=4, value=run.get("remark", ""))
        ws2.cell(row=i, column=5, value=run.get("executed_by", ""))
        ws2.cell(row=i, column=6, value=run.get("executed_at", ""))
    for c in range(1, 7):
        ws2.column_dimensions[chr(64 + c)].width = [8, 30, 10, 20, 10, 20][c - 1]
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("缺陷列表")
    cols3 = ["缺陷ID", "标题", "严重程度", "优先级", "状态", "指派人ID"]
    for c, h in enumerate(cols3, start=1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for i, d in enumerate(data.get("defects", []), start=2):
        ws3.cell(row=i, column=1, value=d.get("id", ""))
        ws3.cell(row=i, column=2, value=d.get("title", ""))
        ws3.cell(row=i, column=3, value=d.get("severity", ""))
        ws3.cell(row=i, column=4, value=d.get("priority", ""))
        ws3.cell(row=i, column=5, value=d.get("status", ""))
        ws3.cell(row=i, column=6, value=d.get("assignee_id", ""))
    for c in range(1, 7):
        ws3.column_dimensions[chr(64 + c)].width = [8, 30, 10, 8, 12, 10][c - 1]
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_op(db, current, "导出", target=f"报告#{r.id}", detail=r.title)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=report_{r.id}.xlsx"},
    )


@router.delete("/{report_id}")
def delete_report(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    r = db.get(TestReport, report_id)
    if not r:
        raise HTTPException(404, "报告不存在")
    if current.role != "admin":
        raise HTTPException(403, "仅管理员可删除报告")
    db.delete(r)
    log_op(db, current, "删除", target=f"报告#{report_id}", detail=r.title)
    db.commit()
    return {"detail": "删除成功"}
