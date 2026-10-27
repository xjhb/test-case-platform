import io

from datetime import datetime

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import accessible_project_ids, get_current_user, require_project_access
from ..logger import log_op
from ..models import TestCase, TestCaseVersion, TestPlan, TestRun, User
from ..schemas import TestCaseCreate, TestCaseOut, TestCaseUpdate

router = APIRouter(prefix="/api/testcases", tags=["测试用例"], dependencies=[Depends(get_current_user)])

VALID_PRIORITY = {"P0", "P1", "P2", "P3"}


def _save_version(db: Session, tc: TestCase, user: User, summary: str = ""):
    latest = db.execute(
        select(TestCaseVersion)
        .where(TestCaseVersion.testcase_id == tc.id)
        .order_by(TestCaseVersion.version.desc())
        .limit(1)
    ).scalar_one_or_none()
    next_ver = (latest.version + 1) if latest else 1
    v = TestCaseVersion(
        testcase_id=tc.id, version=next_ver,
        module=tc.module, title=tc.title,
        precondition=tc.precondition, steps=tc.steps,
        expected=tc.expected, priority=tc.priority,
        changed_by=user.id, change_summary=summary,
    )
    db.add(v)


@router.get("/export")
def export_testcases(
    project_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """导出某项目下所有测试用例为 Excel"""
    require_project_access(project_id, db, current)
    cases = list(
        db.execute(select(TestCase).where(TestCase.project_id == project_id).order_by(TestCase.id.asc())).scalars()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    cols = ["ID", "模块", "用例标题", "优先级", "前置条件", "测试步骤", "预期结果", "当前状态", "创建时间"]
    for c, h in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, tc in enumerate(cases, start=2):
        ws.cell(row=r, column=1, value=tc.id)
        ws.cell(row=r, column=2, value=tc.module)
        ws.cell(row=r, column=3, value=tc.title)
        ws.cell(row=r, column=4, value=tc.priority)
        ws.cell(row=r, column=5, value=tc.precondition)
        ws.cell(row=r, column=6, value=tc.steps)
        ws.cell(row=r, column=7, value=tc.expected)
        ws.cell(row=r, column=8, value=tc.status)
        ws.cell(row=r, column=9, value=tc.created_at.strftime("%Y-%m-%d %H:%M") if tc.created_at else "")
    widths = [6, 14, 30, 8, 24, 30, 24, 12, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_op(db, current, "导出", target=f"用例列表 项目#{project_id}", detail=f"{len(cases)}条")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=testcases.xlsx"},
    )


@router.get("/stats")
def stats(
    project_id: int | None = None,
    plan_id: int | None = None,
    stage: str | None = None,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    accessible = accessible_project_ids(db, current)
    q = select(TestCase)
    if project_id is not None:
        require_project_access(project_id, db, current)
        q = q.where(TestCase.project_id == project_id)
    elif accessible is not None:
        if not accessible:
            return {"total": 0, "by_status": {"未执行": 0, "通过": 0, "失败": 0, "阻塞": 0}, "by_priority": {}}
        q = q.where(TestCase.project_id.in_(accessible))
    cases = list(db.execute(q).scalars())
    total = len(cases)

    by_priority: dict[str, int] = {}
    for c in cases:
        by_priority[c.priority] = by_priority.get(c.priority, 0) + 1

    by_status: dict[str, int] = {"未执行": 0, "通过": 0, "失败": 0, "阻塞": 0}
    if plan_id and stage:
        case_ids = [c.id for c in cases]
        runs = list(
            db.execute(
                select(TestRun)
                .where(TestRun.plan_id == plan_id, TestRun.stage == stage)
                .order_by(TestRun.id.desc())
            ).scalars()
        )
        latest: dict[int, str] = {}
        for r in runs:
            if r.testcase_id not in latest:
                latest[r.testcase_id] = r.status
        for cid in case_ids:
            st = latest.get(cid, "未执行")
            by_status[st] = by_status.get(st, 0) + 1
    else:
        by_status["未执行"] = total

    return {"total": total, "by_status": by_status, "by_priority": by_priority}


@router.get("")
def list_testcases(
    project_id: int | None = None,
    plan_id: int | None = None,
    stage: str | None = None,
    status: str | None = None,
    priority: str | None = None,
    review_status: str | None = None,
    keyword: str | None = None,
    all_stages: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    accessible = accessible_project_ids(db, current)
    q = select(TestCase)
    if project_id is not None:
        require_project_access(project_id, db, current)
        q = q.where(TestCase.project_id == project_id)
    elif accessible is not None:
        if not accessible:
            return {"total": 0, "page": page, "page_size": page_size, "items": []}
        q = q.where(TestCase.project_id.in_(accessible))
    if priority:
        q = q.where(TestCase.priority == priority)
    if review_status:
        q = q.where(TestCase.review_status == review_status)
    if keyword:
        kw = f"%{keyword}%"
        q = q.where(TestCase.title.like(kw) | TestCase.module.like(kw))

    all_rows = list(db.execute(q.order_by(TestCase.id.desc())).scalars())

    STAGES = ["SIT", "UAT", "PRE"]

    # ---- all_stages 模式:超管总览,返回三阶段各自状态(忽略 stage 参数) ----
    if all_stages and current.role == "admin" and plan_id:
        plan = db.get(TestPlan, plan_id)
        cur_stage = plan.current_stage if plan else "SIT"
        stage_maps: dict[str, dict[int, dict]] = {}
        for st in STAGES:
            runs = list(
                db.execute(
                    select(TestRun)
                    .where(TestRun.plan_id == plan_id, TestRun.stage == st)
                    .order_by(TestRun.id.desc())
                ).scalars()
            )
            m: dict[int, dict] = {}
            for r in runs:
                if r.testcase_id not in m:
                    m[r.testcase_id] = {"status": r.status, "screenshot": r.screenshot}
            stage_maps[st] = m

        items_all = []
        for r in all_rows:
            # Manually construct dict to handle None values
            data_dict = {
                "id": r.id,
                "project_id": r.project_id,
                "module": r.module,
                "title": r.title,
                "precondition": r.precondition or "",
                "steps": r.steps or "",
                "expected": r.expected or "",
                "priority": r.priority,
                "status": r.status,
                "review_status": r.review_status or "草稿",
                "reviewer_id": r.reviewer_id,
                "reviewer_name": "",
                "review_comment": r.review_comment or "",
                "reviewed_at": r.reviewed_at,
                "created_by": r.created_by,
                "created_at": r.created_at,
                "updated_at": r.updated_at,
            }
            d = TestCaseOut.model_validate(data_dict).model_dump(mode="json")
            if r.reviewer_id:
                reviewer = db.get(User, r.reviewer_id)
                d["reviewer_name"] = reviewer.display_name or reviewer.username if reviewer else ""
            for st in STAGES:
                info = stage_maps[st].get(r.id)
                d[f"{st.lower()}_status"] = info["status"] if info else "未执行"
                d[f"{st.lower()}_screenshot"] = info["screenshot"] if info else ""
            d["latest_status"] = d.get(f"{cur_stage.lower()}_status", "未执行")
            items_all.append(d)

        if status:
            items_all = [it for it in items_all if any(it.get(f"{s.lower()}_status") == status for s in STAGES)]
        total = len(items_all)
        page_items = items_all[(page - 1) * page_size : page * page_size]
        return {"total": total, "page": page, "page_size": page_size, "items": page_items, "all_stages": True}

    # ---- 普通模式:单阶段状态 ----
    latest_map: dict[int, dict] = {}
    if plan_id and stage:
        runs = list(
            db.execute(
                select(TestRun)
                .where(TestRun.plan_id == plan_id, TestRun.stage == stage)
                .order_by(TestRun.id.desc())
            ).scalars()
        )
        for r in runs:
            if r.testcase_id not in latest_map:
                latest_map[r.testcase_id] = {
                    "latest_status": r.status,
                    "latest_run_id": r.id,
                    "latest_screenshot": r.screenshot,
                }

    items_all = []
    for r in all_rows:
        # Manually construct dict to handle None values
        data_dict = {
            "id": r.id,
            "project_id": r.project_id,
            "module": r.module,
            "title": r.title,
            "precondition": r.precondition or "",
            "steps": r.steps or "",
            "expected": r.expected or "",
            "priority": r.priority,
            "status": r.status,
            "review_status": r.review_status or "草稿",
            "reviewer_id": r.reviewer_id,
            "reviewer_name": "",
            "review_comment": r.review_comment or "",
            "reviewed_at": r.reviewed_at,
            "created_by": r.created_by,
            "created_at": r.created_at,
            "updated_at": r.updated_at,
        }
        d = TestCaseOut.model_validate(data_dict).model_dump(mode="json")
        # 填充评审人姓名
        if r.reviewer_id:
            reviewer = db.get(User, r.reviewer_id)
            d["reviewer_name"] = reviewer.display_name or reviewer.username if reviewer else ""
        lr = latest_map.get(r.id, {})
        d["latest_status"] = lr.get("latest_status", "未执行")
        d["latest_run_id"] = lr.get("latest_run_id")
        d["latest_screenshot"] = lr.get("latest_screenshot", "")
        items_all.append(d)

    if status:
        items_all = [it for it in items_all if it["latest_status"] == status]

    total = len(items_all)
    page_items = items_all[(page - 1) * page_size : page * page_size]
    return {"total": total, "page": page, "page_size": page_size, "items": page_items, "all_stages": False}


@router.post("", response_model=TestCaseOut)
def create_testcase(payload: TestCaseCreate, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    if payload.priority not in VALID_PRIORITY:
        raise HTTPException(status_code=400, detail="优先级只能是 P0/P1/P2/P3")
    require_project_access(payload.project_id, db, current)
    tc = TestCase(
        project_id=payload.project_id,
        module=payload.module,
        title=payload.title,
        precondition=payload.precondition,
        steps=payload.steps,
        expected=payload.expected,
        priority=payload.priority,
        status="未执行",
        created_by=current.id,
    )
    db.add(tc)
    db.commit()
    db.refresh(tc)
    _save_version(db, tc, current, "初始创建")
    db.commit()
    log_op(db, current, "创建", target=f"用例#{tc.id} {tc.title}", detail=f"项目#{payload.project_id}")
    return tc


@router.get("/{tcid}", response_model=TestCaseOut)
def get_testcase(tcid: int, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    return tc


@router.put("/{tcid}", response_model=TestCaseOut)
def update_testcase(tcid: int, payload: TestCaseUpdate, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    data = payload.model_dump(exclude_unset=True)
    if "priority" in data and data["priority"] not in VALID_PRIORITY:
        raise HTTPException(status_code=400, detail="优先级只能是 P0/P1/P2/P3")
    for k, v in data.items():
        setattr(tc, k, v)
    _save_version(db, tc, current, "编辑保存")
    db.commit()
    db.refresh(tc)
    log_op(db, current, "编辑", target=f"用例#{tc.id} {tc.title}")
    return tc


@router.delete("/{tcid}")
def delete_testcase(tcid: int, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    title = tc.title
    db.delete(tc)
    db.commit()
    log_op(db, current, "删除", target=f"用例#{tcid} {title}")
    return {"ok": True}


# ---------- 评审流程 ----------
VALID_REVIEW_STATUS = {"草稿", "评审中", "已批准", "已驳回"}


@router.post("/{tcid}/submit-review")
def submit_review(tcid: int, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    if tc.review_status not in ("草稿", "已驳回"):
        raise HTTPException(status_code=400, detail=f"当前状态[{tc.review_status}]不允许提交评审,只有草稿或被驳回的用例可以提交评审")
    tc.review_status = "评审中"
    tc.reviewer_id = None
    tc.review_comment = ""
    tc.reviewed_at = None
    db.commit()
    db.refresh(tc)
    log_op(db, current, "提交评审", target=f"用例#{tc.id} {tc.title}")
    return tc


@router.post("/{tcid}/approve")
def approve_testcase(
    tcid: int,
    comment: str = Body("", embed=True),
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    if tc.review_status != "评审中":
        raise HTTPException(status_code=400, detail="只有评审中的用例才能被批准")
    tc.review_status = "已批准"
    tc.reviewer_id = current.id
    tc.review_comment = comment
    tc.reviewed_at = datetime.now()
    db.commit()
    db.refresh(tc)
    log_op(db, current, "批准", target=f"用例#{tc.id} {tc.title}")
    return tc


@router.post("/{tcid}/reject")
def reject_testcase(
    tcid: int,
    body: dict = Body(...),
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    tc = db.get(TestCase, tcid)
    if not tc:
        raise HTTPException(status_code=404, detail="用例不存在")
    require_project_access(tc.project_id, db, current)
    if tc.review_status != "评审中":
        raise HTTPException(status_code=400, detail="只有评审中的用例才能被驳回")
    comment = body.get("comment", "")
    if not comment.strip():
        raise HTTPException(status_code=400, detail="驳回时必须填写意见")
    tc.review_status = "已驳回"
    tc.reviewer_id = current.id
    tc.review_comment = comment
    tc.reviewed_at = datetime.now()
    db.commit()
    db.refresh(tc)
    log_op(db, current, "驳回", target=f"用例#{tc.id} {tc.title}", detail=comment[:50])
    return tc
